import openpyxl
import os

FILE_NAME = "student_results.xlsx"


def create_excel_file():
    if not os.path.exists(FILE_NAME):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "Student Results"

        headers = [
            "Roll No",
            "Name",
            "Class",
            "OS",
            "CP",
            "TCO",
            "ML",
            "ES",
            "Total",
            "Percentage",
            "Grade",
            "Status"
        ]

        sheet.append(headers)
        workbook.save(FILE_NAME)
        workbook.close()


def calculate_result(marks):
    total = sum(marks)
    percentage = total / 5

    if all(mark >= 35 for mark in marks):
        status = "PASS"

        if percentage >= 90:
            grade = "A+"
        elif percentage >= 80:
            grade = "A"
        elif percentage >= 70:
            grade = "B"
        elif percentage >= 60:
            grade = "C"
        elif percentage >= 50:
            grade = "D"
        else:
            grade = "E"
    else:
        grade = "F"
        status = "FAIL"

    return total, percentage, grade, status


def add_student():
    workbook = openpyxl.load_workbook(FILE_NAME)
    sheet = workbook.active

    try:
        roll_no = int(input("Enter Roll No: "))
    except ValueError:
        print("Enter a valid Roll No.")
        workbook.close()
        return

    for row in sheet.iter_rows(min_row=2, values_only=True):
        if row[0] == roll_no:
            print("Roll No already exists!")
            workbook.close()
            return

    name = input("Enter Student Name: ")
    student_class = input("Enter Class/Course: ")

    marks = []

    subjects = ["OS", "CP", "TCO", "ML", "ES"]

    for subject in subjects:
        while True:
            try:
                mark = float(input(f"Enter Marks for {subject}: "))

                if 0 <= mark <= 100:
                    marks.append(mark)
                    break
                else:
                    print("Marks must be between 0 and 100.")

            except ValueError:
                print("Enter a valid number.")

    total, percentage, grade, status = calculate_result(marks)

    sheet.append([
        roll_no,
        name,
        student_class,
        marks[0],
        marks[1],
        marks[2],
        marks[3],
        marks[4],
        total,
        percentage,
        grade,
        status
    ])

    workbook.save(FILE_NAME)
    workbook.close()

    print("Student result added successfully.")
    print("Roll No:", roll_no)
    print("Student Name:", name)
    print("Class:", student_class)
    print("Total Marks:", total)
    print("Percentage:", f"{percentage:.2f}%")
    print("Grade:", grade)
    print("Result Status:", status)


def get_result():
    workbook = openpyxl.load_workbook(FILE_NAME)
    sheet = workbook.active

    try:
        roll_no = int(input("Enter Roll No: "))
    except ValueError:
        print("Enter a valid Roll No.")
        workbook.close()
        return

    found = False

    for row in sheet.iter_rows(min_row=2, values_only=True):
        if row[0] == roll_no:
            print("Roll No:", row[0])
            print("Name:", row[1])
            print("Class:", row[2])
            print("OS:", row[3])
            print("CP:", row[4])
            print("TCO:", row[5])
            print("ML:", row[6])
            print("ES:", row[7])
            print("Total:", row[8])
            print("Percentage:", f"{row[9]:.2f}%")
            print("Grade:", row[10])
            print("Status:", row[11])

            found = True
            break

    workbook.close()

    if not found:
        print("Student record not found.")


def show_all_data():
    workbook = openpyxl.load_workbook(FILE_NAME)
    sheet = workbook.active

    data_found = False

    for row in sheet.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue

        data_found = True

        print("Roll No:", row[0])
        print("Name:", row[1])
        print("Class:", row[2])
        print("Total:", row[8])
        print("Percentage:", f"{row[9]:.2f}%")
        print("Grade:", row[10])
        print("Status:", row[11])
        print()

    workbook.close()

    if not data_found:
        print("Student records are not available.")


def menu():
    create_excel_file()

    while True:
        print("1. Add Student Result")
        print("2. Get Student Result")
        print("3. Show All Student Results")
        print("4. Exit")

        choice = input("Enter your choice: ")

        if choice == "1":
            add_student()

        elif choice == "2":
            get_result()

        elif choice == "3":
            show_all_data()

        elif choice == "4":
            print("Thank you for visiting")
            break

        else:
            print("Invalid choice")


menu()